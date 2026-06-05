import yfinance as yf
import pandas as pd
import requests
import json
import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

# ==========================================
# 1. الإعدادات
# ==========================================
TELEGRAM_BOT_TOKEN = os.environ.get('TELEGRAM_BOT_TOKEN', '')
TELEGRAM_CHAT_ID = os.environ.get('TELEGRAM_CHAT_ID', '')
ALPACA_API_KEY = os.environ.get('ALPACA_API_KEY', '')
ALPACA_SECRET_KEY = os.environ.get('ALPACA_SECRET_KEY', '')

MODEL_FILE = "lstm_model.json"

# قوائم المراقبة
MARKETS = {
    "الأسهم الأمريكية": ["AAPL", "TSLA", "NVDA", "MSFT", "AMZN", "SPY"],
    "الأسهم السعودية": ["2222.SR", "1120.SR"],
    "الأسهم الإماراتية": ["ADCB.AD", "FAB.AD"],
    "العملات الرقمية": ["BTC-USD", "ETH-USD"],
    "المعادن": ["GC=F", "SI=F"],
    "الفوركس": ["EURUSD=X", "GBPUSD=X", "USDJPY=X", "AUDUSD=X", "USDCAD=X"]
}

# ==========================================
# 2. دوال مساعدة
# ==========================================
def send_telegram(message):
    if not TELEGRAM_BOT_TOKEN or not TELEGRAM_CHAT_ID:
        return False
    try:
        url = f"https://api.telegram.org/bot{TELEGRAM_BOT_TOKEN}/sendMessage"
        payload = {"chat_id": TELEGRAM_CHAT_ID, "text": message, "parse_mode": "HTML"}
        r = requests.post(url, json=payload, timeout=10)
        return r.status_code == 200
    except:
        return False

def send_telegram_document(file_path):
    """إرسال ملف عبر تيليجرام"""
    if not TELEGRAM_BOT_TOKEN or not TELEGRAM_CHAT_ID:
        return False
    try:
        url = f"https://api.telegram.org/bot{TELEGRAM_BOT_TOKEN}/sendDocument"
        with open(file_path, 'rb') as f:
            files = {'document': f}
            data = {'chat_id': TELEGRAM_CHAT_ID}
            r = requests.post(url, data=data, files=files, timeout=30)
        return r.status_code == 200
    except Exception as e:
        print(f"خطأ إرسال الملف: {e}")
        return False

def calculate_rsi(df, period=14):
    try:
        delta = df['Close'].diff()
        gain = (delta.where(delta > 0, 0)).rolling(period).mean()
        loss = (-delta.where(delta < 0, 0)).rolling(period).mean()
        rs = gain / loss
        return 100 - (100 / (1 + rs))
    except:
        return None

def analyze_ticker(ticker):
    """تحليل شامل للسهم"""
    try:
        df = yf.Ticker(ticker).history(period="60d", interval="1d")
        if df.empty or len(df) < 20:
            return None
        
        current_price = float(df['Close'].iloc[-1])
        prev_close = float(df['Close'].iloc[-2]) if len(df) > 1 else current_price
        change = current_price - prev_close
        change_pct = (change / prev_close) * 100
        
        # الدعم والمقاومة
        support = float(df['Low'].rolling(20).min().iloc[-1])
        resistance = float(df['High'].rolling(20).max().iloc[-1])
        
        # RSI
        rsi_series = calculate_rsi(df)
        current_rsi = float(rsi_series.iloc[-1]) if rsi_series is not None else 0
        
        # المتوسطات المتحركة
        ma_20 = float(df['Close'].rolling(20).mean().iloc[-1])
        ma_50 = float(df['Close'].rolling(50).mean().iloc[-1]) if len(df) >= 50 else ma_20
        
        # الحجم
        avg_volume = float(df['Volume'].rolling(20).mean().iloc[-1])
        current_volume = float(df['Volume'].iloc[-1])
        volume_ratio = current_volume / avg_volume if avg_volume > 0 else 1
        
        # المسافات
        dist_support = ((current_price - support) / current_price) * 100
        dist_resistance = ((resistance - current_price) / current_price) * 100
        
        # تحديد الإشارة
        signal = "مراقبة"
        signal_strength = "ضعيفة"
        recommendation = "انتظار"
        
        if dist_support <= 2.5 and current_rsi < 35:
            signal = "شراء قوي"
            signal_strength = "قوية جداً"
            recommendation = "شراء"
        elif dist_support <= 4.0 and current_rsi < 40:
            signal = "شراء"
            signal_strength = "متوسطة"
            recommendation = "شراء حذر"
        elif dist_resistance <= 2.5 and current_rsi > 65:
            signal = "بيع قوي"
            signal_strength = "قوية جداً"
            recommendation = "بيع"
        elif dist_resistance <= 4.0 and current_rsi > 60:
            signal = "بيع"
            signal_strength = "متوسطة"
            recommendation = "بيع حذر"
        elif current_rsi < 30:
            signal = "تشبع بيعي"
            signal_strength = "متوسطة"
            recommendation = "مراقبة للشراء"
        elif current_rsi > 70:
            signal = "تشبع شرائي"
            signal_strength = "متوسطة"
            recommendation = "مراقبة للبيع"
        
        return {
            'ticker': ticker,
            'price': current_price,
            'change': change,
            'change_pct': change_pct,
            'rsi': current_rsi,
            'support': support,
            'resistance': resistance,
            'dist_support': dist_support,
            'dist_resistance': dist_resistance,
            'ma_20': ma_20,
            'ma_50': ma_50,
            'volume': current_volume,
            'volume_ratio': volume_ratio,
            'signal': signal,
            'signal_strength': signal_strength,
            'recommendation': recommendation
        }
    except Exception as e:
        print(f"خطأ في {ticker}: {e}")
        return None

def load_lstm_models():
    """تحميل نماذج LSTM"""
    if not os.path.exists(MODEL_FILE):
        return {}
    try:
        with open(MODEL_FILE, 'r') as f:
            return json.load(f)
    except:
        return {}

# ==========================================
# 3. إنشاء تقرير Excel
# ==========================================
def create_excel_report():
    """إنشاء تقرير Excel شامل"""
    print(f" إنشاء تقرير Excel - {datetime.now()}")
    
    wb = Workbook()
    wb.remove(wb.active)  # إزالة sheet الافتراضي
    
    # ألوان التنسيق
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    
    buy_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    sell_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    watch_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    lstm_models = load_lstm_models()
    all_data = {}
    
    for market_name, tickers in MARKETS.items():
        print(f"📈 تحليل {market_name}...")
        market_data = []
        
        for ticker in tickers:
            result = analyze_ticker(ticker)
            if result:
                # إضافة تأكيد LSTM إذا متاح
                lstm_confirmed = "غير متاح"
                if ticker in lstm_models:
                    lstm_confirmed = "متاح"
                
                result['lstm_available'] = lstm_confirmed
                market_data.append(result)
        
        all_data[market_name] = market_data
        
        # إنشاء sheet لكل سوق
        ws = wb.create_sheet(title=market_name[:31])  # Excel يقبل 31 حرف كحد أقصى
        
        # العناوين
        headers = ['الرمز', 'السعر', 'التغيير', 'التغيير %', 'RSI', 
                   'الدعم', 'المقاومة', 'المسافة من الدعم %', 'المسافة من المقاومة %',
                   'MA 20', 'MA 50', 'الحجم', 'نسبة الحجم', 
                   'الإشارة', 'قوة الإشارة', 'التوصية', 'LSTM']
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
            cell.border = thin_border
        
        # البيانات
        for row_idx, data in enumerate(market_data, 2):
            ws.cell(row=row_idx, column=1, value=data['ticker']).border = thin_border
            ws.cell(row=row_idx, column=2, value=round(data['price'], 2)).border = thin_border
            ws.cell(row=row_idx, column=3, value=round(data['change'], 2)).border = thin_border
            ws.cell(row=row_idx, column=4, value=round(data['change_pct'], 2)).border = thin_border
            ws.cell(row=row_idx, column=5, value=round(data['rsi'], 1)).border = thin_border
            ws.cell(row=row_idx, column=6, value=round(data['support'], 2)).border = thin_border
            ws.cell(row=row_idx, column=7, value=round(data['resistance'], 2)).border = thin_border
            ws.cell(row=row_idx, column=8, value=round(data['dist_support'], 2)).border = thin_border
            ws.cell(row=row_idx, column=9, value=round(data['dist_resistance'], 2)).border = thin_border
            ws.cell(row=row_idx, column=10, value=round(data['ma_20'], 2)).border = thin_border
            ws.cell(row=row_idx, column=11, value=round(data['ma_50'], 2)).border = thin_border
            ws.cell(row=row_idx, column=12, value=int(data['volume'])).border = thin_border
            ws.cell(row=row_idx, column=13, value=round(data['volume_ratio'], 2)).border = thin_border
            ws.cell(row=row_idx, column=14, value=data['signal']).border = thin_border
            ws.cell(row=row_idx, column=15, value=data['signal_strength']).border = thin_border
            ws.cell(row=row_idx, column=16, value=data['recommendation']).border = thin_border
            ws.cell(row=row_idx, column=17, value=data['lstm_available']).border = thin_border
            
            # تلوين حسب الإشارة
            signal_cell = ws.cell(row=row_idx, column=14)
            if 'شراء' in data['signal']:
                signal_cell.fill = buy_fill
            elif 'بيع' in data['signal']:
                signal_cell.fill = sell_fill
            else:
                signal_cell.fill = watch_fill
            
            # تلوين التغيير
            change_cell = ws.cell(row=row_idx, column=4)
            if data['change_pct'] > 0:
                change_cell.font = Font(color="006100")
            elif data['change_pct'] < 0:
                change_cell.font = Font(color="9C0006")
        
        # تعديل عرض الأعمدة
        for col in range(1, 18):
            ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 15
    
    # Sheet الملخص العام
    summary_ws = wb.create_sheet(title="الملخص العام", index=0)
    
    summary_ws['A1'] = "تقرير المراقبة الشامل"
    summary_ws['A1'].font = Font(bold=True, size=16, color="366092")
    summary_ws['A2'] = f"التاريخ: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    summary_ws['A2'].font = Font(size=11)
    
    summary_ws['A4'] = "السوق"
    summary_ws['B4'] = "عدد الأسهم"
    summary_ws['C4'] = "إشارات شراء"
    summary_ws['D4'] = "إشارات بيع"
    summary_ws['E4'] = "مراقبة"
    
    for col in range(1, 6):
        cell = summary_ws.cell(row=4, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
    
    row = 5
    total_buy = 0
    total_sell = 0
    total_watch = 0
    
    for market_name, data in all_data.items():
        buy_count = sum(1 for d in data if 'شراء' in d['signal'])
        sell_count = sum(1 for d in data if 'بيع' in d['signal'])
        watch_count = len(data) - buy_count - sell_count
        
        total_buy += buy_count
        total_sell += sell_count
        total_watch += watch_count
        
        summary_ws.cell(row=row, column=1, value=market_name).border = thin_border
        summary_ws.cell(row=row, column=2, value=len(data)).border = thin_border
        summary_ws.cell(row=row, column=3, value=buy_count).border = thin_border
        summary_ws.cell(row=row, column=4, value=sell_count).border = thin_border
        summary_ws.cell(row=row, column=5, value=watch_count).border = thin_border
        row += 1
    
    # الإجمالي
    summary_ws.cell(row=row, column=1, value="الإجمالي").font = Font(bold=True)
    summary_ws.cell(row=row, column=2, value=sum(len(d) for d in all_data.values())).font = Font(bold=True)
    summary_ws.cell(row=row, column=3, value=total_buy).font = Font(bold=True, color="006100")
    summary_ws.cell(row=row, column=4, value=total_sell).font = Font(bold=True, color="9C0006")
    summary_ws.cell(row=row, column=5, value=total_watch).font = Font(bold=True)
    
    for col in range(1, 6):
        summary_ws.cell(row=row, column=col).border = thin_border
    
    # Sheet التوصيات
    rec_ws = wb.create_sheet(title="التوصيات")
    
    rec_ws['A1'] = "التوصيات الاستثمارية"
    rec_ws['A1'].font = Font(bold=True, size=16, color="366092")
    
    rec_ws['A3'] = "الرمز"
    rec_ws['B3'] = "السعر"
    rec_ws['C3'] = "الإشارة"
    rec_ws['D3'] = "التوصية"
    rec_ws['E3'] = "السبب"
    
    for col in range(1, 6):
        cell = rec_ws.cell(row=3, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
    
    row = 4
    for market_name, data in all_data.items():
        for d in data:
            if d['signal'] != "مراقبة":
                reason = f"RSI: {d['rsi']:.1f}, المسافة من الدعم: {d['dist_support']:.1f}%"
                rec_ws.cell(row=row, column=1, value=d['ticker']).border = thin_border
                rec_ws.cell(row=row, column=2, value=d['price']).border = thin_border
                rec_ws.cell(row=row, column=3, value=d['signal']).border = thin_border
                rec_ws.cell(row=row, column=4, value=d['recommendation']).border = thin_border
                rec_ws.cell(row=row, column=5, value=reason).border = thin_border
                row += 1
    
    # حفظ الملف
    filename = f"تقرير_المراقبة_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    wb.save(filename)
    print(f"✅ تم حفظ التقرير: {filename}")
    
    return filename, {
        'total': sum(len(d) for d in all_data.values()),
        'buy': total_buy,
        'sell': total_sell,
        'watch': total_watch
    }

# ==========================================
# 4. التشغيل
# ==========================================
def main():
    print(f"📊 بدء إنشاء تقرير Excel - {datetime.now()}")
    send_telegram(f"📊 <b>جاري إنشاء تقرير Excel الشامل...</b>\n⏰ {datetime.now().strftime('%H:%M')}")
    
    try:
        filename, stats = create_excel_report()
        
        # إرسال الملخص
        summary = f"""
✅ <b>تم إنشاء التقرير بنجاح!</b>

 <b>الإحصائيات:</b>
• إجمالي الأصول: {stats['total']}
• إشارات شراء: {stats['buy']}
• إشارات بيع: {stats['sell']}
• مراقبة: {stats['watch']}

📁 <b>اسم الملف:</b> {filename}

 {datetime.now().strftime('%Y-%m-%d %H:%M')}
"""
        send_telegram(summary)
        
        # إرسال الملف
        print(f" جاري إرسال الملف عبر تيليجرام...")
        if send_telegram_document(filename):
            print("✅ تم إرسال الملف بنجاح")
            send_telegram("📎 <b>تم إرسال ملف Excel بنجاح!</b>")
        else:
            print("⚠️ فشل إرسال الملف")
            send_telegram("⚠️ <b>فشل إرسال الملف. يمكنك تحميله من GitHub.</b>")
        
    except Exception as e:
        print(f"❌ خطأ: {e}")
        send_telegram(f"❌ <b>فشل إنشاء التقرير:</b>\n{str(e)[:200]}")

if __name__ == "__main__":
    main()
